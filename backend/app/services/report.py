import os, re
from xml.sax.saxutils import escape
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                Spacer, PageBreak)
from reportlab.platypus.tableofcontents import TableOfContents
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def _clean(t: str) -> str:
    return re.sub(r"[*#`]+", "", t or "").strip()

def _paras(t: str):
    return [p.strip() for p in _clean(t).split("\n\n") if p.strip()]

def _sections(rep):
    secs = [("Executive Summary", rep.get("executive_summary", ""))]
    secs += [(s["heading"], s["content"]) for s in rep.get("sections", [])]
    secs.append(("Conclusion", rep.get("conclusion", "")))
    return secs

# ---------------- PDF ----------------
class _Doc(BaseDocTemplate):
    def __init__(self, fn, **kw):
        super().__init__(fn, pagesize=A4, **kw)
        f = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")
        self.addPageTemplates([PageTemplate(id="p", frames=f, onPage=self._footer)])

    def _footer(self, canv, doc):
        canv.setFont("Helvetica", 9)
        canv.drawCentredString(A4[0] / 2, 1.2 * cm, f"Page {doc.page}")

    def afterFlowable(self, fl):
        if isinstance(fl, Paragraph) and fl.style.name == "H1":
            self.notify("TOCEntry", (0, fl.getPlainText(), self.page))

def make_pdf(rep: dict, path: str):
    ss = getSampleStyleSheet()
    title = ParagraphStyle("T", parent=ss["Title"], fontSize=26, spaceAfter=20)
    h1 = ParagraphStyle("H1", parent=ss["Heading1"], spaceBefore=14, spaceAfter=8)
    body = ParagraphStyle("B", parent=ss["BodyText"], leading=16, spaceAfter=8)
    small = ParagraphStyle("S", parent=ss["BodyText"], fontSize=9, leading=12)

    story = [Spacer(1, 6 * cm), Paragraph(escape(_clean(rep["title"])), title),
             Paragraph("AI-Generated Research Report", ss["Heading3"]), PageBreak(),
             Paragraph("Table of Contents", ss["Heading1"])]
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("L0", fontSize=12, leading=18)]
    story += [toc, PageBreak()]
    for head, content in _sections(rep):
        story.append(Paragraph(escape(_clean(head)), h1))
        story += [Paragraph(escape(p), body) for p in _paras(content)]
    story += [Paragraph("References", h1)]
    for s in rep["sources"]:
        story.append(Paragraph(f"[{s['id']}] {escape(s['title'])} - {escape(s['url'])}", small))
    _Doc(path, leftMargin=2.2*cm, rightMargin=2.2*cm, topMargin=2.2*cm,
         bottomMargin=2.2*cm).multiBuild(story)

# ---------------- DOCX ----------------
def _add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for typ, txt in [("begin", None), (None, 'TOC \\o "1-1" \\h \\z \\u'), ("separate", None)]:
        if typ:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
    t = OxmlElement("w:t"); t.text = "Right-click here and choose 'Update Field' to build the table of contents."
    r._r.append(t)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
    # ask Word to refresh fields on open
    s = OxmlElement("w:updateFields"); s.set(qn("w:val"), "true")
    doc.settings.element.append(s)

def make_docx(rep: dict, path: str):
    d = Document()
    d.add_heading(_clean(rep["title"]), 0)
    d.add_paragraph("AI-Generated Research Report")
    d.add_page_break()
    d.add_heading("Table of Contents", 1)
    _add_toc(d)
    d.add_page_break()
    for head, content in _sections(rep):
        d.add_heading(_clean(head), 1)
        for p in _paras(content):
            d.add_paragraph(p)
    d.add_heading("References", 1)
    for s in rep["sources"]:
        d.add_paragraph(f"[{s['id']}] {s['title']} - {s['url']}")
    d.save(path)

def generate_files(job_id: str, rep: dict, out_dir: str):
    os.makedirs(out_dir, exist_ok=True)
    make_pdf(rep, os.path.join(out_dir, f"{job_id}.pdf"))
    make_docx(rep, os.path.join(out_dir, f"{job_id}.docx"))
